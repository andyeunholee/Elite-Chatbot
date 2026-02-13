import streamlit as st
import asyncio
from agent import agent
from langchain_core.messages import HumanMessage, AIMessage, SystemMessage

# Page Config
st.set_page_config(page_title="Elite U.S. College Advisor", page_icon="🎓")

# Display Logo and Title
col1, col2 = st.columns([1, 5])
with col1:
    st.image("elitelogo-Blue-removebg-preview.png", width=80)
with col2:
    st.title("Elite U.S. College Advisor")

# Removed caption as requested
# st.caption("Powered by Google Gemini & Tavily Search")

# Initialize Session State
if "messages" not in st.session_state:
    st.session_state["messages"] = [
        SystemMessage(content="You are an elite US college admissions AI consultant named 'Genny', serving students nationwide (California, Georgia, etc.). Your role is to provide comprehensive, up-to-date news and strategic advice on US university admissions. Always use the search tool to find the latest specific local data when asked. IMPORTANT: Answer in the SAME language as the user's question. If the user asks in Korean, answer in Korean. If in English, answer in English."),
        AIMessage(content="Hello! I am **Genny**, an AI Agent specializing in US college admissions consulting.\n\nAsk me anything!")
    ]

# Display Chat History
for msg in st.session_state.messages:
    if isinstance(msg, HumanMessage):
        with st.chat_message("user"):
            st.markdown(msg.content)
    elif isinstance(msg, AIMessage):
        with st.chat_message("assistant"):
            # Handle list content for history display
            content_to_display = msg.content
            if isinstance(content_to_display, list):
                text_content = ""
                for block in content_to_display:
                    if isinstance(block, dict) and block.get("type") == "text":
                        text_content += block.get("text", "")
                content_to_display = text_content
            st.markdown(content_to_display)

# Handle User Input
if prompt := st.chat_input("Ask about college chances, advise, or news..."):
    # Add user message to state
    st.session_state.messages.append(HumanMessage(content=prompt))
    with st.chat_message("user"):
        st.markdown(prompt)

    # Process with Agent
    with st.chat_message("assistant"):
        message_placeholder = st.empty()
        full_response = ""
        
        # Create a container for status updates (tool calls)
        status_container = st.status("Thinking...", expanded=False)
        
        try:
            # Optimize inputs for the graph
            input_messages = st.session_state.messages
            
            # Stream the graph execution
            # stream_mode="updates" allows us to see intermediate steps (tool calls)
            for event in agent.stream({"messages": input_messages}, stream_mode="updates"):
                for node_name, node_output in event.items():
                    if node_name == "chatbot":
                        # The LLM generated a message (either a tool call or a final answer)
                        message = node_output["messages"][-1]
                        if message.tool_calls:
                            status_container.update(label="Searching the web...", state="running", expanded=True)
                            for tool_call in message.tool_calls:
                                status_container.write(f"🔍 Searching: {tool_call['args']}")
                        else:
                            # Final answer (or intermediate thought)
                            # Handle both string and list content (Gemini/Claude sometimes returns list of blocks)
                            if isinstance(message.content, list):
                                full_response = ""
                                for block in message.content:
                                    if isinstance(block, dict) and block.get("type") == "text":
                                        full_response += block.get("text", "")
                            else:
                                full_response = message.content
                    
                    elif node_name == "tools":
                        # The tool executed
                        status_container.update(label="Reading results...", state="running", expanded=True)
                        tool_message = node_output["messages"][-1]
                        status_container.write("✅ Found results.")
            
            status_container.update(label="Finished!", state="complete", expanded=False)
            message_placeholder.markdown(full_response)
            
            # Save assistant response to state
            st.session_state.messages.append(AIMessage(content=full_response))
            
        except Exception as e:
            st.error(f"An error occurred: {e}") 

# Sidebar for Student Profile Upload
with st.sidebar:
    st.header("📂 Student Profile (Upload)")
    st.write("Upload transcripts, resume, or SAT scores.")
    uploaded_files = st.file_uploader("Upload Profile Data", type=["pdf", "txt", "docx", "xlsx", "csv", "jpg", "png", "jpeg"], accept_multiple_files=True)
    
    # Store student data in session state to persist it during interaction
    if "student_data" not in st.session_state:
        st.session_state["student_data"] = ""
    
    if uploaded_files:
        temp_data = ""
        processing_msg = st.status("Processing files...", expanded=True)
        
        for uploaded_file in uploaded_files:
            try:
                processing_msg.write(f"Reading {uploaded_file.name}...")
                
                # PDF
                if uploaded_file.type == "application/pdf":
                    import pypdf
                    reader = pypdf.PdfReader(uploaded_file)
                    for page in reader.pages:
                        temp_data += page.extract_text() + "\n"
                
                # Text
                elif uploaded_file.type == "text/plain":
                    temp_data += uploaded_file.read().decode("utf-8") + "\n"
                
                # Word (.docx)
                elif uploaded_file.name.endswith(".docx"):
                    import docx
                    doc = docx.Document(uploaded_file)
                    for para in doc.paragraphs:
                        temp_data += para.text + "\n"
                
                # Excel/CSV (.xlsx, .csv)
                elif uploaded_file.name.endswith(".xlsx") or uploaded_file.name.endswith(".csv"):
                    import pandas as pd
                    if uploaded_file.name.endswith(".csv"):
                        df = pd.read_csv(uploaded_file)
                    else:
                        df = pd.read_excel(uploaded_file)
                    temp_data += df.to_string() + "\n"
                
                # Images (JPG, PNG) - Using Gemini Vision for OCR
                elif uploaded_file.type.startswith("image/"):
                    from PIL import Image
                    import google.generativeai as genai
                    
                    # Open image
                    image = Image.open(uploaded_file)
                    
                    # Use a lightweight Gemini model just for text extraction
                    vision_model = genai.GenerativeModel('gemini-2.0-flash') 
                    response = vision_model.generate_content(["Transcribe all text from this image exactly as it appears.", image])
                    temp_data += f"\n[Image Content from {uploaded_file.name}]:\n{response.text}\n"

            except Exception as e:
                st.error(f"Error reading {uploaded_file.name}: {e}")
        
        if temp_data:
            st.session_state["student_data"] += temp_data # Append new file content
            processing_msg.update(label="Profile Loaded! ✅", state="complete", expanded=False)
            st.success("New File Loaded! ✅")

    # --- Google Sheets Database (Shared) ---
    st.markdown("---")
    st.header("☁️ Student Database (Google Sheets)")
    
    # Google Sheets Connection
    import gspread
    from oauth2client.service_account import ServiceAccountCredentials
    import os

    # Scope for Google Sheets and Drive API
    # Scope for Google Sheets and Drive API
    scope = ['https://spreadsheets.google.com/feeds', 'https://www.googleapis.com/auth/drive']
    creds_file = 'google_secret.json'
    
    creds = None
    
    # 1. Try Loading from Streamlit Secrets (Cloud)
    if "gcp_service_account" in st.secrets:
        try:
            # Create a dictionary from secrets
            service_account_info = st.secrets["gcp_service_account"]
            creds = ServiceAccountCredentials.from_json_keyfile_dict(service_account_info, scope)
        except Exception as e:
            st.error(f"Error loading secrets: {e}")

    # 2. Try Loading from Local File (Local Dev)
    elif os.path.exists(creds_file):
        try:
            creds = ServiceAccountCredentials.from_json_keyfile_name(creds_file, scope)
        except Exception as e:
            st.error(f"Error loading local file: {e}")
            
    # Check if credentials were loaded
    if not creds:
        st.error("⚠️ Google Cloud Credentials not found! Please check secrets or upload 'google_secret.json'.")
    else:
        try:
            client = gspread.authorize(creds)
            
            # Connect to the Sheet
            SHEET_NAME = "Elite_Student_DB"
            try:
                sheet = client.open(SHEET_NAME).sheet1
            except gspread.exceptions.SpreadsheetNotFound:
                st.error(f"⚠️ Sheet '{SHEET_NAME}' not found! Please create a Google Sheet named '{SHEET_NAME}' and share it with the service account email.")
                sheet = None

            if sheet:
                # SAVE functionality
                student_name_input = st.text_input("Enter Student Name to Save")
                if st.button("Save to Cloud DB"):
                    if student_name_input and st.session_state["student_data"]:
                        # Check if student already exists
                        try:
                            # Search for the student name in the first column
                            cell = sheet.find(student_name_input, in_column=1)
                            if cell:
                                # Update existing row (Name is col 1, Data is col 2)
                                sheet.update_cell(cell.row, 2, st.session_state["student_data"])
                                st.success(f"Updated profile for {student_name_input}!")
                            else:
                                # Append new row
                                sheet.append_row([student_name_input, st.session_state["student_data"]])
                                st.success(f"Saved new profile for {student_name_input}!")
                        except Exception as e:
                             # Fallback if find fails or other API error
                            st.error(f"Error saving to Google Sheets: {e}")

                    elif not st.session_state["student_data"]:
                        st.warning("No student data to save yet. Upload a file first.")
                    else:
                        st.warning("Please enter a name.")

                # LOAD functionality
                try:
                    # Get all records (list of dicts)
                    records = sheet.get_all_records()
                    # Extract names (assuming header is 'Student Name')
                    # If header is missing or different, we fallback to column 1 values
                    if records and "Student Name" in records[0]:
                        saved_students = [r["Student Name"] for r in records if r["Student Name"]]
                    else:
                        # Fallback: get all values from column 1, skipping header
                        col1_values = sheet.col_values(1)
                        saved_students = col1_values[1:] if len(col1_values) > 1 else []
                    
                    selected_student = st.selectbox("Load Saved Student", ["Select a student..."] + saved_students)
                    
                    if selected_student != "Select a student...":
                        if st.button("Load Profile"):
                            # Find the row for the selected student
                            # We search again to be robust
                             cell = sheet.find(selected_student, in_column=1)
                             if cell:
                                 # Get data from column 2
                                 loaded_data = sheet.cell(cell.row, 2).value
                                 st.session_state["student_data"] = loaded_data
                                 st.success(f"Loaded profile: {selected_student}")
                             else:
                                 st.error("Student not found in DB.")
                
                except Exception as e:
                    st.error(f"Error loading from Google Sheets: {e}")

        except Exception as e:
            st.error(f"Error connecting to Google Sheets: {e}")

    # --- Inject into System Prompt ---
    # Only if we have student data in session state
    if st.session_state["student_data"]:
        system_msg_content = st.session_state["messages"][0].content
        
         # Reset student data part if it already exists to allow updates
        base_system_msg = system_msg_content.split("\n\n[CURRENT STUDENT DATA]:")[0]
        
        st.session_state["messages"][0] = SystemMessage(content=base_system_msg + f"\n\n[CURRENT STUDENT DATA]:\n{st.session_state['student_data']}\n\nINSTRUCTION: The user has uploaded the above student profile. Use this data to provide personalized admission prediction ('Chance Me') and specific improvement advice when asked.")
        
        with st.expander("Review Current Analyzed Data"):
            # Limit display to avoid UI clutter
            display_text = st.session_state["student_data"][:500] + "..." if len(st.session_state["student_data"]) > 500 else st.session_state["student_data"]
            st.text(display_text)

    # --- Clear / Reset ---
    st.markdown("---")
    if st.button("🧹 Clear Profile / New Chat"):
        # Clear student data
        st.session_state["student_data"] = ""
        # Reset chat history to initial state
        st.session_state["messages"] = [
            SystemMessage(content="You are an elite US college admissions AI consultant named 'Genny', serving students nationwide (California, Georgia, etc.). Your role is to provide comprehensive, up-to-date news and strategic advice on US university admissions. Always use the search tool to find the latest specific local data when asked. IMPORTANT: Answer in the SAME language as the user's question. If the user asks in Korean, answer in Korean. If in English, answer in English."),
            AIMessage(content="Hello! I am **Genny**, an AI Agent specializing in US college admissions consulting.\n\nAsk me anything!")
        ]
        st.rerun()

